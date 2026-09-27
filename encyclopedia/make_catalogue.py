"""Regenerate src/tool_catalogue.json from Part D of the (corrected) workbook."""
import json, re
from openpyxl import load_workbook
ws = load_workbook("The_Encyclopedia.xlsx")["Reference"]
cat = {}
for r in range(507, 651):
    for c in range(17, ws.max_column + 1):
        v = ws.cell(r, c).value
        if v and "AGE:" in str(v):
            v = str(v).replace("\n\n*** FULL TEACHING — PART G ***", "")
            name = v.split("\n")[0].strip()
            f = {"name": name}
            for key in ("AGE", "MEASURES", "CANNOT TELL YOU", "TIME", "LANDS IN"):
                m = re.search(key + r": (.*)", v)
                if m: f[key] = m.group(1).strip()
            if name not in cat or len(v) > len(cat[name]["_raw"]):
                f["_raw"] = v; cat[name] = f
json.dump(cat, open("src/tool_catalogue.json", "w"), indent=1, ensure_ascii=False)
print("catalogue: %d tools" % len(cat))
