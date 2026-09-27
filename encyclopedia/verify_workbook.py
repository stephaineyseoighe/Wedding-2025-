"""Integrity checks on the built workbook. Exit 1 on any failure."""
import sys, statistics
from openpyxl import load_workbook

book = sys.argv[1] if len(sys.argv) > 1 else "The_Encyclopedia.xlsx"
base = load_workbook("src/base.xlsx")
wb = load_workbook(book)
fail = []

if wb.sheetnames != base.sheetnames: fail.append("sheet list changed")
# Excel's hard cell limit
over = [(s, c.coordinate, len(c.value)) for s in wb.sheetnames for row in wb[s].iter_rows() for c in row
        if isinstance(c.value, str) and len(c.value) > 32767]
if over: fail.append("cells over 32,767 chars: %s" % over[:10])
longest = max(((len(c.value), s, c.coordinate) for s in wb.sheetnames for row in wb[s].iter_rows() for c in row
               if isinstance(c.value, str)), default=(0, "", ""))
# untouched sheets identical
for s in ("Log", "Appendix 4 and 5", "Cavan plan", "Sources", "Start here", "Lists"):
    a, b = base[s], wb[s]
    d = sum(1 for row in a.iter_rows() for c in row if c.value != b[c.coordinate].value
            and not (s == "Appendix 4 and 5" and c.column == 16))  # column P: worked examples
    if d: fail.append("%s changed in %d cells" % (s, d))
    if len(a.data_validations.dataValidation) != len(b.data_validations.dataValidation): fail.append("%s dropdowns changed" % s)
# formulas preserved (Reference K/L count columns point at Log)
fa = [(c.coordinate, c.value) for row in base["Reference"].iter_rows(max_row=669) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
fb = [(c.coordinate, c.value) for row in wb["Reference"].iter_rows(max_row=669) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
if fa != fb: fail.append("Reference formulas changed (%d vs %d)" % (len(fa), len(fb)))
# internal hyperlinks still land on a non-empty row
for row in wb["Reference"].iter_rows():
    for c in row:
        if c.hyperlink and c.hyperlink.location:
            tgt = c.hyperlink.location.split("!")[-1].replace("$", "")
            r = int("".join(ch for ch in tgt if ch.isdigit()))
            if not any(wb["Reference"].cell(r, k).value for k in range(1, 6)): fail.append("dead link %s → %s" % (c.coordinate, tgt))
# error values
errs = [(s, c.coordinate) for s in wb.sheetnames for row in wb[s].iter_rows() for c in row
        if isinstance(c.value, str) and c.value.strip() in ("#REF!", "#VALUE!", "#NAME?", "#N/A", "#DIV/0!", "nan")]
if errs: fail.append("error-like values: %s" % errs[:10])
# content counts
sp = wb["Skill progression"]
mic = [i for i in range(2, sp.max_row + 1) if sp.cell(i, 1).value == "MICRO-SKILL"]
rich = [i for i in mic if not str(sp.cell(i, 7).value).startswith("NOT YET")]
L = [len(str(sp.cell(i, c).value or "")) for i in rich for c in range(8, 13)]
ref = wb["Reference"]
kinds = {}
for r in range(670, ref.max_row + 1):
    k = ref.cell(r, 1).value
    if k in ("TOOL", "METHOD", "CONDITION", "PRESENTATION"): kinds[k] = kinds.get(k, 0) + 1
cond = sum(1 for r in range(1, wb["Conditions"].max_row + 1) if wb["Conditions"].cell(r, 1).value == "CONDITION")
print("micro-skills %d/%d · median stage cell %d" % (len(rich), len(mic), statistics.median(L)))
print("Part G tools %s · Part H methods %s · Part I conditions %s · presentations %s" % (
    kinds.get("TOOL"), kinds.get("METHOD"), kinds.get("CONDITION"), kinds.get("PRESENTATION")))
print("Conditions sheet: %d condition rows (%d conditions × 5 bands) · %d rows" % (cond, cond // 5, wb["Conditions"].max_row))
ap = wb["Appendix 4 and 5"]
ex = sum(1 for r in range(2, ap.max_row + 1) if ap.cell(r, 16).value and "NOT FOR SUBMISSION" in str(ap.cell(r, 16).value))
print("worked examples beside reflection prompts: %d" % ex)
print("longest cell: %d chars (%s %s)" % longest)
if len(rich) != len(mic): fail.append("unwritten micro-skills remain")
print("\n".join(fail) or "ALL CHECKS PASSED")
sys.exit(1 if fail else 0)
