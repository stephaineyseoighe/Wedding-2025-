"""Write micro-skills from micro_new.py into the Skill progression sheet.

Stands in for progression.py when the original build folder is not to hand.
Matches rows on (macro skill, micro-skill), writes columns G–N, P and U,
leaves formatting, stage-tracking column O and target dates untouched.

    python3 apply_micro.py            # writes into The_Encyclopedia.xlsx
"""
import statistics
import sys
from openpyxl import load_workbook

from micro_new import MICRO

BOOK = "The_Encyclopedia.xlsx"


def norm(s):
    return " ".join(str(s or "").split())


wb = load_workbook(BOOK)
ws = wb["Skill progression"]
M = ws.max_row
index = {}
for i in range(2, M + 1):
    if ws.cell(i, 1).value == "MICRO-SKILL":
        index[(norm(ws.cell(i, 3).value), norm(ws.cell(i, 5).value))] = i

problems = []
for (macro, micro), d in MICRO.items():
    r = index.get((norm(macro), norm(micro)))
    if r is None:
        problems.append("no row for %r / %r" % (macro, micro))
        continue
    old4 = str(ws.cell(r, 11).value or "")
    if "PASS TEST — " in old4 and not old4.startswith("YOU DO IT ALONE"):
        pt = old4.split("PASS TEST — ")[-1].strip()
        if "PASS TEST — " + pt not in d["s4"]:
            problems.append("pass test not verbatim in %r: %r" % (micro, pt))
            continue
    ws.cell(r, 7).value = d["why"]
    for c, k in zip(range(8, 13), ("s1", "s2", "s3", "s4", "s5")):
        ws.cell(r, c).value = d[k]
    ws.cell(r, 13).value = d["error"]
    ws.cell(r, 14).value = d["feeds"]
    ws.cell(r, 16).value = d["s1"].split("\n\n")[0]
    ws.cell(r, 21).value = d["evidence"]

if problems:
    print("\n".join(problems))
    sys.exit(1)

wb.save(BOOK)

mic = [i for i in range(2, M + 1) if ws.cell(i, 1).value == "MICRO-SKILL"]
rich = [i for i in mic if not str(ws.cell(i, 7).value).startswith("NOT YET")]
L = [len(str(ws.cell(i, c).value or "")) for i in rich for c in range(8, 13)]
new = [len(d[k]) for d in MICRO.values() for k in ("s1", "s2", "s3", "s4", "s5")]
print("written: %d of %d | median stage cell: %d | median of new cells: %d"
      % (len(rich), len(mic), statistics.median(L), statistics.median(new)))
