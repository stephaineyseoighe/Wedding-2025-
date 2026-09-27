"""Rebuild the Conditions sheet: existing blocks kept, new conditions from records/cond_*.py
appended inside each of the five age bands (condition row, co-occurring rows, presentation row).

Run on a workbook that has not had it applied (build_all.sh starts from src/base.xlsx).
"""
import json
from openpyxl import load_workbook
from reclib import (load, snapshot, clear, put, make, BANDS, BAND_FULL, BAND_CAPS, BAND_TABLE3)

BOOK = "The_Encyclopedia.xlsx"
wb = load_workbook(BOOK)
ws = wb["Conditions"]
NC = ws.max_column
CAT = json.load(open("src/tool_catalogue.json"))
V = "▸  "


def sec(head, items):
    return head + "\n" + "\n\n".join(V + x.strip() for x in items)


# ── snapshot existing sheet, grouped by band ──────────────────────────────────
rows = [snapshot(ws, r, NC) for r in range(1, ws.max_row + 1)]
level = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
top = []
groups = {}          # band full name -> [snapshots]; the AGE BAND row first
cur = None
for s, lv in zip(rows, level):
    if lv == "AGE BAND":
        cur = s["cells"][2][0]; groups[cur] = [s]
    elif cur is None:
        top.append(s)
    else:
        groups[cur].append(s)
T_COND = next(s for s, lv in zip(rows, level) if lv == "CONDITION")
T_CO = next(s for s, lv in zip(rows, level) if lv == "CO-OCCURRING")
T_PRES = next(s for s, lv in zip(rows, level) if lv == "PRESENTATION")
existing = {s["cells"][4][0] for s, lv in zip(rows, level) if lv == "CONDITION"}
PRES_GENERIC = {c: T_PRES["cells"][c - 1][0] for c in range(10, NC + 1)}   # generic presentation text


def tool_text(name):
    if name in CAT:
        t = CAT[name]
        out = "%s — AGE %s" % (t["name"], t.get("AGE", "check the manual"))
        for k in ("MEASURES", "CANNOT TELL YOU", "TIME"):
            if t.get(k):
                out += "\n     %s: %s" % (k, t[k])
        return out
    return name.replace(" · ", "\n     ")


def short(name):
    return name.split(" (")[0] if len(name) > 40 else name


conds = [c for c in load("CONDS", "cond_*.py") if c["name"] not in existing]
for c in conds:
    b, pw = c["bands"], c["pathway"]
    diagnosed = "NOT A DSM DIAGNOSIS" if c["code"].lower().startswith("not a dsm") else "DIAGNOSED"
    co = ["%s — %s. PRESENTS AS: %s" % (x["name"].upper(), x["rate"].rstrip("."), x["presents"]) for x in c["cooccurring"]]
    where = sec("WHERE THIS SITS", ["NEPS referral form category: " + c["neps"], "UCD competency: 1. Assessment",
                                    "CORU: " + c["coru"], "PSI: " + c["psi"], "Law and policy: " + c["law"]])
    shared = {
        9: sec("WHAT IT IS, AND WHAT IT IS NOT", c["what_it_is"] + c["what_it_is_not"]),
        11: "THE PATHWAY AND THE TIMELINE\n" + "\n\n".join(
            "▸ %s\n   %s" % (h, pw[k]) for h, k in (("TYPICAL AGE OF IDENTIFICATION, AND WHY THEN", "age"),
                                                  ("WHO DIAGNOSES IT", "who_diagnoses"),
                                                  ("WHO COULD HAVE WRITTEN THE REPORT IN FRONT OF YOU", "who_wrote_report"),
                                                  ("WHO YOU REFER TO", "refer_to"), ("'SHOULD WE HAVE COME SOONER?'", "sooner"))),
        13: sec("RED FLAGS AND BOUNDARIES", c["red_flags"]),
        14: sec("WHAT ELSE COULD IT BE, AND WHAT CO-OCCURS", c["differential"] + co),
        16: sec("WHAT YOU DO, AND WHAT YOU WRITE", c["recommendations"]),
        17: "EXPLAIN\n" + "\n\n".join(
            [V + "TO A PARENT"] + [V + x for x in c["explain_parent"]] + [V + "TO A TEACHER"] + [V + x for x in c["explain_teacher"]]
            + [V + "TO THE CHILD"] + [V + x for x in c["explain_child"]] + [V + "ANALOGIES"] + [V + x for x in c["analogies"]]
            + [V + "LANGUAGE"] + [V + x for x in c["language"]]),
        18: sec("THE CHILD’S VOICE, AND THE QUESTIONS", c["child_voice"] + c["questions"]),
        19: "YOUR PRACTICE\n" + "\n\n".join([V + "SUPERVISION"] + [V + x for x in c["supervision"]]
                                           + [V + "REFLECTION"] + [V + x for x in c["reflection"]]),
        20: sec("WHAT TO DO NEXT", c["next"]),
        22: sec("CITATIONS", c["citations"]),
    }
    for band in BANDS:
        x = b[band]
        other = " · ".join(BAND_CAPS[k] for k in BANDS if k != band)
        vals = {1: "CONDITION", 2: "1. Assessment", 3: BAND_FULL[band], 4: diagnosed, 5: c["name"], 6: c["code"],
                7: x["applies"], 8: where,
                10: sec("HOW COMMON AT THIS AGE", c["prevalence"] + ["%s: %s" % (BAND_CAPS[band], x["prevalence"]),
                                                                    "AT OTHER AGES (for context): " + other]),
                12: sec("WHAT YOU SEE AT THIS AGE", [x["see"]] + ["%s: %s" % (BAND_CAPS[k], b[k]["see"]) for k in BANDS]),
                15: sec("TOOLS VALID AT THIS BAND", [tool_text(t) for t in x["tools"]] or
                        ["No standardised tool is the starting point at this band — history, observation and the referral route carry it. See WHAT YOU SEE."]),
                21: sec("WHERE THE EVIDENCE GOES", ["Psychological report", "Appendix 5 entry tagged 1. Assessment",
                                                    "Appendix 7 if observed", "Table 3 — " + BAND_TABLE3[band],
                                                    "Log sheet — tag the referral area and the competency"])}
        vals.update(shared)
        new = [make(T_COND, vals)]
        for item in c["cooccurring"]:
            new.append(make(T_CO, {
                1: "CO-OCCURRING", 2: "1. Assessment", 3: BAND_FULL[band], 4: "CO-OCCURS WITH " + short(c["name"]).upper(),
                5: item["name"].upper(), 6: "see its own entry", 7: "YES — assess separately",
                8: sec("WHERE THIS SITS", ["Sits under 1. Assessment, alongside %s." % short(c["name"]), "NEPS category: " + c["neps"]]),
                9: sec("HOW OFTEN IT CO-OCCURS", [item["rate"]]),
                10: sec("HOW IT PRESENTS WHEN IT CO-OCCURS", [item["presents"]]),
                11: sec("THE RISK", ["Do not assume the primary diagnosis explains this.",
                                     "Write down before you start what the primary diagnosis does NOT explain, then assess that.",
                                     "Diagnostic overshadowing is the single most common error in co-occurring presentations."]),
                12: "—", 13: "—", 14: "—",
                15: sec("ASSESS", ["Assess this separately, with its own measures.",
                                   "Build the case for each condition independently rather than letting evidence for one carry the other."]),
                16: sec("DO", ["State in the report what would be needed to establish this second presentation.",
                               "That is often a referral rather than a conclusion you can reach yourself."]),
                17: "—", 18: "—",
                19: sec("YOUR PRACTICE", ["Bring it to supervision: for each child on my caseload with this diagnosis, what has NOT been assessed because the diagnosis was already there?"]),
                20: sec("WHAT TO DO NEXT", ["Assess separately, then refer if the threshold is met."]),
                21: sec("WHERE THE EVIDENCE GOES", ["Psychological report — state the co-occurrence and what it changes."]),
                22: "—"}))
        pv = dict(PRES_GENERIC)
        pv.update({1: "PRESENTATION", 2: "1. Assessment", 3: BAND_FULL[band], 4: "SYMPTOM-ONLY",
                   5: "Presentations related to %s — no diagnosis" % short(c["name"]), 6: "No DSM code",
                   7: "YES — these present at every band",
                   8: sec("WHERE THIS SITS", ["NEPS category: as above", "UCD competency: 1. Assessment",
                                              "CORU: " + c["coru"], "PSI: " + c["psi"]]),
                   9: sec("THE PRESENTATIONS — no diagnosis needed, and most referrals look like this", c["presentations"])})
        new.append(make(T_PRES, pv))
        groups[BAND_FULL[band]].extend(new)

# ── write back ────────────────────────────────────────────────────────────────
n = len(existing) + len(conds)
v, st, hl = top[1]["cells"][4]
top[1]["cells"][4] = ("%d conditions · 5 age bands" % n, st, hl)
clear(ws, 1, NC)
r = 1
for s in top:
    put(ws, r, s); r += 1
for band in [BAND_FULL[k] for k in BANDS]:
    for s in groups[band]:
        put(ws, r, s); r += 1
wb.save(BOOK)
print("Conditions sheet: +%d conditions (%d total) · %d rows" % (len(conds), n, r - 1))
