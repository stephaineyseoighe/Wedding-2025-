"""Add two trial sheets showing the two ways of making one all-in-one sheet.

Option 1 — the merged Conditions sheet (Conditions layout), copied under a new name,
           with every internal link re-pointed at the copy.
Option 2 — Skill progression exactly as it is (stage layout), with an EASY READ column,
           and below it every other row of the merged Conditions sheet (conditions,
           tools, standards, papers …) under its own header row.

The original sheets are left untouched: the Log's formulas and dropdowns still use them.
"""
import copy, re
from openpyxl import load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from reclib import snapshot, put, make

BOOK = "The_Encyclopedia.xlsx"
OPT1, OPT2 = "Option 1 – All in one", "Option 2 – Skill layout"
SKILL_KINDS = {"COMPETENCY", "MACRO SKILL", "MICRO-SKILL"}

wb = load_workbook(BOOK)
C, S = wb["Conditions"], wb["Skill progression"]
NC = 23
norm = lambda s: " ".join(str(s or "").split()).lower()


def repoint(cell, sheet, rowmap=None):
    """Move an internal 'Conditions'!A<n> link onto `sheet` (row n → rowmap[n] if given)."""
    h = cell.hyperlink
    if not h or not h.location or not h.location.startswith("'Conditions'!"):
        return
    n = int(re.search(r"A(\d+)$", h.location).group(1))
    if rowmap is not None:
        if n not in rowmap:
            return  # target not on this sheet: keep the link to Conditions
        n = rowmap[n]
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location="'%s'!A%d" % (sheet, n))


for name in (OPT1, OPT2):
    if name in wb.sheetnames:
        del wb[name]

# ── Option 1 ─────────────────────────────────────────────────────────────────
o1 = wb.copy_worksheet(C)
o1.title = OPT1
o1.freeze_panes = C.freeze_panes
o1.sheet_properties.tabColor = "1F6F5F"
for row in o1.iter_rows():
    for cell in row:
        repoint(cell, OPT1)

# ── Option 2 ─────────────────────────────────────────────────────────────────
o2 = wb.copy_worksheet(S)
o2.title = OPT2
o2.freeze_panes = S.freeze_panes
o2.sheet_properties.tabColor = "8A5A00"
SC = S.max_column
EC = SC + 1  # Easy Read column after Skill progression's own columns

# Skill rows: find each Conditions skill row's twin in Skill progression
sp_rows = {}
for r in range(2, S.max_row + 1):
    k = S.cell(r, 1).value
    if k in SKILL_KINDS:
        for c in (5, 3, 2):
            v = norm(S.cell(r, c).value)
            if v and r not in sp_rows.setdefault((k, v), []):
                sp_rows[(k, v)].append(r)
rowmap, easy_n = {}, 0
for n in range(2, C.max_row + 1):
    k = C.cell(n, 1).value
    if k not in SKILL_KINDS:
        continue
    cands = next((sp_rows[(k, norm(C.cell(n, c).value))] for c in (5, 3, 2) if sp_rows.get((k, norm(C.cell(n, c).value)))), None)
    if not cands:
        continue
    r = cands.pop(0)  # names that repeat are matched in order
    rowmap[n] = r
    if C.cell(n, 23).value and not o2.cell(r, EC).value:
        o2.cell(r, EC).value = C.cell(n, 23).value
        o2.cell(r, EC)._style = copy.copy(S.cell(r, SC)._style)
        easy_n += 1
o2.cell(1, EC).value = "EASY READ — in plain words"
o2.cell(1, EC)._style = copy.copy(S.cell(1, SC)._style)
o2.column_dimensions[o2.cell(1, EC).column_letter].width = 60

# Everything else from the merged sheet, below the skill table
start = S.max_row + 2
banner = snapshot(C, 2, NC)
out = start
put(o2, out, make(banner, {1: "SECTION", 5: "EVERYTHING ELSE — conditions, tools, age bands, presentations, methods, "
                                          "standards, papers. From here down the columns follow the Conditions sheet; "
                                          "the header row below names them."}, height=45))
out += 1
put(o2, out, snapshot(C, 1, NC)); out += 1
moved = []
for n in range(2, C.max_row + 1):
    if C.cell(n, 1).value in SKILL_KINDS:
        continue
    put(o2, out, snapshot(C, n, NC))
    rowmap[n] = out
    moved.append(out)
    out += 1
# widen to the Conditions columns where Skill progression's are narrower
for c in range(1, NC + 1):
    L = o2.cell(1, c).column_letter
    w1, w2 = o2.column_dimensions[L].width, C.column_dimensions[L].width
    if w2 and (not w1 or w2 > w1):
        o2.column_dimensions[L].width = w2
for row in o2.iter_rows(min_row=start):
    for cell in row:
        repoint(cell, OPT2, rowmap)
wb.save(BOOK)
print("options: %s %d rows · %s %d rows (skill table %d rows + %d rows below, Easy Read on %d skill rows)"
      % (OPT1, o1.max_row, OPT2, o2.max_row, S.max_row, len(moved), easy_n))
