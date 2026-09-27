"""Validate a parts/*.py file against the workbook and the HANDOVER standard.

    python3 check_part.py parts/intervention_a.py
"""
import importlib.util, statistics, sys
from openpyxl import load_workbook

path = sys.argv[1]
spec = importlib.util.spec_from_file_location("part", path)
mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
MICRO = mod.MICRO
ws = load_workbook("The_Encyclopedia.xlsx")["Skill progression"]
norm = lambda s: " ".join(str(s or "").split())
rows = {(norm(ws.cell(i, 3).value), norm(ws.cell(i, 5).value)): i
        for i in range(2, ws.max_row + 1) if ws.cell(i, 1).value == "MICRO-SKILL"}
LABELS = {"s1": ("WATCH FOR — ", "WRITE DOWN — ", "ASK AFTERWARDS — "),
          "s2": ("YOU TAKE — ", "THEY HOLD — ", "WHY THAT SPLIT — "),
          "s3": ("YOU RUN IT — ", "WHAT THEY ARE WATCHING FOR — ", "STEPPING IN LOOKS LIKE — "),
          "s4": ("YOU DO IT ALONE — ", "PASS TEST — ", "EVIDENCE TO FILE — ", "WHAT FAILURE LOOKS LIKE — "),
          "s5": ("YOU CAN TEACH IT — ", "WHEN IT GOES WRONG — ")}
bad, lens = [], []
for (macro, micro), d in MICRO.items():
    r = rows.get((norm(macro), norm(micro)))
    if not r: bad.append("NO ROW: %s / %s" % (macro, micro)); continue
    for k in ("why", "s1", "s2", "s3", "s4", "s5", "error", "feeds", "evidence"):
        if not d.get(k): bad.append("MISSING %s: %s" % (k, micro))
    pt = str(ws.cell(r, 11).value or "").split("PASS TEST — ")[-1].strip()
    if "PASS TEST — " + pt not in d.get("s4", ""): bad.append("PASS TEST not verbatim: %s -> %r" % (micro, pt))
    for k, labs in LABELS.items():
        for lab in labs:
            if lab not in d.get(k, ""): bad.append("LABEL %r missing in %s: %s" % (lab, k, micro))
        n = len(d.get(k, "")); lens.append(n)
        if n < 250: bad.append("THIN %s (%d chars): %s" % (k, n, micro))
    if not d.get("feeds", "").startswith("→ Feeds into"): bad.append("feeds must start '→ Feeds into': %s" % micro)
print("\n".join(bad) or "OK")
print("entries: %d | median stage cell: %s" % (len(MICRO), statistics.median(lens) if lens else "-"))
sys.exit(1 if bad else 0)
