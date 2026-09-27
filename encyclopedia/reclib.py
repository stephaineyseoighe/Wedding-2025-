"""Shared helpers for the Part G/H/I and Conditions builders."""
import copy, glob, importlib.util, json, re

BANDS = ["Early Years", "School Age", "Adolescent", "Young Adult", "Special Setting"]
BAND_FULL = {"Early Years": "Early Years / Pre-School (0–5 years)", "School Age": "School Age (6–12 years)",
             "Adolescent": "Adolescent (13–16 years)", "Young Adult": "Young Adult (17–26 years)",
             "Special Setting": "Special Setting (class or school)"}
BAND_CAPS = {"Early Years": "EARLY YEARS 0–5", "School Age": "SCHOOL AGE 6–12", "Adolescent": "ADOLESCENT 13–16",
             "Young Adult": "YOUNG ADULT 17–26", "Special Setting": "SPECIAL SETTING"}
BAND_TABLE3 = {"Early Years": "Early Years", "School Age": "School Age", "Adolescent": "Adolescent",
               "Young Adult": "Young Adult", "Special Setting": "the special setting row"}
URL = re.compile(r"https?://\S+")


def load(var, pattern):
    """All records of one kind, from records/<pattern>, in file order."""
    out = []
    for p in sorted(glob.glob("records/" + pattern)):
        s = importlib.util.spec_from_file_location(p, p)
        m = importlib.util.module_from_spec(s); s.loader.exec_module(m)
        out.extend(getattr(m, var))
    return out


def dots(items, mark="•  "):
    return "\n\n".join(mark + x.strip() for x in items)


def arrows(items):
    return "\n\n".join("→ " + x.strip() for x in items)


def first_url(s):
    m = URL.search(s or "")
    return m.group(0).rstrip(").,;") if m else None


def snapshot(ws, r, ncol):
    """Everything needed to re-create row r elsewhere."""
    cells = []
    for c in range(1, ncol + 1):
        cell = ws.cell(r, c)
        hl = (cell.hyperlink.target, cell.hyperlink.location) if cell.hyperlink else None
        cells.append((cell.value, copy.copy(cell._style), hl))
    d = ws.row_dimensions[r]
    return {"cells": cells, "height": d.height, "level": d.outline_level}


def clear(ws, r0, ncol):
    for r in range(r0, ws.max_row + 1):
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            cell.value = None; cell.hyperlink = None
            cell._style = copy.copy(ws.cell(ws.max_row + 5, c)._style)
        ws.row_dimensions[r].height = None
        ws.row_dimensions[r].outline_level = 0


def put(ws, r, snap):
    for c, (v, st, hl) in enumerate(snap["cells"], start=1):
        cell = ws.cell(r, c)
        cell.value = v; cell._style = copy.copy(st)
        cell.hyperlink = None
        if hl:
            from openpyxl.worksheet.hyperlink import Hyperlink
            cell.hyperlink = Hyperlink(ref=cell.coordinate, target=hl[0], location=hl[1])
    ws.row_dimensions[r].height = snap["height"]
    ws.row_dimensions[r].outline_level = snap["level"]


def make(template, values, height=None, links=None):
    """New row snapshot: template styles, given {col: value}, optional {col: url}."""
    cells = [(values.get(c), st, None) for c, (_, st, _) in enumerate(template["cells"], start=1)]
    snap = {"cells": cells, "height": height or template["height"], "level": template["level"]}
    for c, url in (links or {}).items():
        v, st, _ = snap["cells"][c - 1]
        snap["cells"][c - 1] = (v, st, (url, None))
    return snap


def co_line(x):
    """One co-occurring condition as 'NAME — rate. PRESENTS AS: …', tolerant of writers repeating the label."""
    pres = re.sub(r"^\s*PRESENTS AS:\s*", "", x["presents"].strip(), flags=re.I)
    return "%s — %s. PRESENTS AS: %s" % (x["name"].upper(), x["rate"].strip().rstrip("."), pres)
