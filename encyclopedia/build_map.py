"""Competency map — one sheet ordered the way the trainee thinks about a case:

    COMPETENCY → AGE BAND → TOPIC (diagnosed + not diagnosed, together) → the skills, beginner to expert

Each competency opens with the skills that apply to every topic (listed once). Then, per age band, every
topic that has skills specific to this competency: the diagnosis (short, with a link to its full entry),
the presentations that look like it without a diagnosis, the tools (Assessment only) and the skills,
in teaching-timetable order, each with a live 1–5 stage bar from Skill progression column O.

Rows are grouped (Excel outline): click + / − at the left, or the 1 2 3 4 buttons at the top left.
Mappings come from records/map_pres.py and records/map_skills_*.py (checked by check_map.py).
"""
import glob, importlib.util, re
from openpyxl import load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.properties import Outline

BOOK = "The_Encyclopedia.xlsx"
NAME = "Competency map"
SP = "'Skill progression'"


def _load(pattern, var):
    out = {}
    for p in sorted(glob.glob("records/" + pattern)):
        s = importlib.util.spec_from_file_location(p, p); m = importlib.util.module_from_spec(s); s.loader.exec_module(m)
        out.update(getattr(m, var))
    return out


PRES_TOPICS = _load("map_pres.py", "PRES_TOPICS")
SKILL_TOPICS = _load("map_skills_*.py", "SKILL_TOPICS")
EASY = _load("easy_*.py", "EASY")

PHASES = ["Before placement", "Weeks 1–2", "Weeks 3–6", "Weeks 7–12", "Ongoing all year"]
PHASE_SHORT = ["Before placement", "Weeks 1–2", "Weeks 3–6", "Weeks 7–12", "All year"]
BANDS = ["Early Years / Pre-School (0–5 years)", "School Age (6–12 years)", "Adolescent (13–16 years)",
         "Young Adult (17–26 years)", "Special Setting (class or school)"]
BAR = ["DCEFEA", "B7DFD5", "86C5B6", "4E9E8C", "1F6F5F"]
INK, DARK, MUTED, LINK = "1B2B28", "123F37", "5B6B67", "1F4E79"
fill = lambda h: PatternFill("solid", fgColor=h)
thin = Side(style="thin", color="C9DCD7")
box = Border(bottom=thin)
wrap = lambda h="left", v="top": Alignment(horizontal=h, vertical=v, wrap_text=True)
F = lambda sz=9, b=False, c=INK, i=False, u=None: Font(name="Arial", size=sz, bold=b, color=c, italic=i, underline=u)

wb = load_workbook(BOOK)
S, C, J = wb["Skill progression"], wb["Conditions"], wb["Skill journey"]

# ── skills, in timetable order, with their Skill journey slide ───────────────
slide_of = {}
for r in range(1, J.max_row + 1):
    h = J.cell(r, 2).hyperlink
    if h and h.location and h.location.startswith(SP + "!O"):
        slide_of[int(h.location.rsplit("O", 1)[1])] = r - 8          # the slide's nav row
skills, phase = [], 4
for r in range(2, S.max_row + 1):
    k = S.cell(r, 1).value
    if k == "BLOCK":
        phase = next((i for i, p in enumerate(PHASES) if p in str(S.cell(r, 6).value or "")), 4)
    elif k == "MICRO-SKILL":
        key = "%s || %s" % (S.cell(r, 3).value, S.cell(r, 5).value)
        skills.append({"row": r, "phase": phase, "comp": S.cell(r, 2).value, "name": S.cell(r, 5).value,
                       "macro": S.cell(r, 3).value, "topics": SKILL_TOPICS.get(key, [])})
skills.sort(key=lambda s: (s["phase"], s["row"]))
comps = list(dict.fromkeys(S.cell(r, 2).value for r in range(2, S.max_row + 1) if S.cell(r, 1).value == "COMPETENCY"))

# ── conditions by band (row in Conditions), presentations (row in Conditions) ─
cond_row, cond_order, band_pres_row = {}, {b: [] for b in BANDS}, {}
pres_row = {}
for r in range(2, C.max_row + 1):
    k, band, name = C.cell(r, 1).value, C.cell(r, 3).value, C.cell(r, 5).value
    if k == "CONDITION" and band in cond_order:
        cond_row[(band, name)] = r; cond_order[band].append(name)
    elif k == "PRESENTATION" and band == "All bands":
        pres_row[name] = r
    elif k == "PRESENTATION" and band in cond_order:
        m = re.match(r"Presentations related to (.+) — no diagnosis", str(name))
        if m: band_pres_row[(band, m.group(1))] = r
pres_by_topic = {}
for p, ts in PRES_TOPICS.items():
    for t in ts: pres_by_topic.setdefault(t, []).append(p)
unlinked = [p for p, ts in PRES_TOPICS.items() if not ts]


def first_points(text, n=2, cap=260):
    """The first n bullet points of a heading-led cell, joined, trimmed."""
    pts = [x.strip(" ▸•→\n") for x in re.split(r"\n\s*(?:▸|•|→)\s*", str(text or ""))[1:] if x.strip(" ▸•→\n")]
    s = " · ".join(pts[:n]) or str(text or "").split("\n", 1)[-1]
    return s if len(s) <= cap else s[:cap].rsplit(" ", 1)[0] + " …"


def easy_first(key, cap=240):
    t = EASY.get(key, "")
    body = [x.strip("• ") for x in t.split("\n")[1:] if x.startswith("•")]
    s = " ".join(body[:2])
    return s if len(s) <= cap else s[:cap].rsplit(" ", 1)[0] + " …"


def tool_lines(text, n=4):
    pts = [x.strip(" ▸•\n") for x in re.split(r"\n\s*▸\s*", str(text or ""))[1:] if x.strip(" ▸•\n")]
    return " · ".join(p.split(" — ")[0][:60] for p in pts[:n])


# ── sheet ────────────────────────────────────────────────────────────────────
if NAME in wb.sheetnames:
    del wb[NAME]
M = wb.create_sheet(NAME, index=wb.sheetnames.index("Conditions"))
M.sheet_view.showGridLines = False
M.sheet_properties.tabColor = "123F37"
M.sheet_properties.outlinePr = Outline(summaryBelow=False, summaryRight=False)
for col, w in {"A": 16, "B": 46, "C": 70, "D": 3.6, "E": 3.6, "F": 3.6, "G": 3.6, "H": 3.6, "I": 20, "J": 15, "K": 16}.items():
    M.column_dimensions[col].width = w
r = 0
here_rules = []


def row(values, level=0, height=None, fl=None, font=None, merge=None, links=None, hidden=False):
    """Write one row. values: {col: value}. merge: (from_col, to_col)."""
    global r
    r += 1
    for c, v in values.items():
        cell = M.cell(r, c); cell.value = v
        cell.font = font or F(); cell.alignment = wrap("left", "center" if (height or 0) <= 20 else "top")
        if fl: cell.fill = fl
    if fl:
        for c in range(1, 12):
            M.cell(r, c).fill = fl
    if merge:
        M.merge_cells(start_row=r, start_column=merge[0], end_row=r, end_column=merge[1])
    for c, loc in (links or {}).items():
        if loc: M.cell(r, c).hyperlink = Hyperlink(ref=M.cell(r, c).coordinate, location=loc)
    d = M.row_dimensions[r]
    if height: d.height = height
    if level: d.outline_level = level
    if hidden: d.hidden = True
    return r


cloc = lambda n: "'Conditions'!A%d" % n
jloc = lambda n: "'Skill journey'!B%d" % n if n else None


def skill_row(s, level, hidden):
    o = "%s!$O$%d" % (SP, s["row"])
    rr = row({1: "SKILL", 2: s["name"], 3: s["macro"],
              **{4 + k: str(k + 1) for k in range(5)},
              9: "=%s" % o, 10: PHASE_SHORT[s["phase"]], 11: "open slide →"},
             level=level, hidden=hidden, height=26, links={2: jloc(slide_of.get(s["row"])), 11: jloc(slide_of.get(s["row"]))})
    M.cell(rr, 1).font = F(8, True, "4E9E8C")
    M.cell(rr, 2).font = F(9, False, LINK, u="single")
    M.cell(rr, 3).font = F(8, False, MUTED, i=True)
    M.cell(rr, 11).font = F(8, False, LINK, u="single")
    for k in range(5):
        c = M.cell(rr, 4 + k); c.alignment = wrap("center", "center"); c.font = F(7, False, "9FB5B0"); c.fill = fill("F2F6F5")
        here_rules.append(("%s%d" % ("DEFGH"[k], rr), o, k))
    M.cell(rr, 9).font = F(8, True, DARK)
    M.cell(rr, 10).font = F(8, False, MUTED)


# title and how to use
row({1: "COMPETENCY MAP — competency › age band › topic › the skills you need, beginner to expert"},
    height=32, fl=fill(DARK), font=F(16, True, "FFFFFF"), merge=(1, 11))
row({1: "HOW TO USE\n"
        "• Rows are grouped. Click + and − on the far left to open and close a section, or the small 1 2 3 4 buttons at the top left "
        "(1 = competencies only, 2 = + age bands, 3 = + topics, 4 = everything).\n"
        "• Under each topic, DIAGNOSED and NOT DIAGNOSED sit together — everything about autism is under Autism, whether or not a child has the diagnosis.\n"
        "• The skills are in the order you learn them (before placement → weeks 1–2 → 3–6 → 7–12 → all year). The 1–5 bar fills in from "
        "Skill progression column O. Click a skill to open its slide in Skill journey.\n"
        "• Blue underlined text is a link to the full entry."},
    height=96, fl=fill("F4FAF8"), font=F(9), merge=(1, 11))
index_at = r + 1
n_index = 1 + len(comps) + (1 if any(not v for v in PRES_TOPICS.values()) else 0) + 5 * sum(
    1 for c in comps if any(s["comp"] == c and s["topics"] and s["topics"] != ["ALL"] for s in skills))
r += n_index + 1  # room for the contents (filled in at the end)
index = []

for comp in comps:
    cs = [s for s in skills if s["comp"] == comp]
    every = [s for s in cs if s["topics"] == ["ALL"] or not s["topics"]]
    specific = [s for s in cs if s["topics"] and s["topics"] != ["ALL"]]
    cr = row({1: comp.upper(), 3: "%d skills · %d for every topic · %d for particular topics" % (len(cs), len(every), len(specific))},
             height=30, fl=fill(DARK), font=F(14, True, "FFFFFF"), merge=(1, 2))
    M.cell(cr, 3).font = F(9, False, "DCEFEA")
    index.append((comp, cr))
    ce = EASY.get("competency::" + comp)
    if ce:
        row({1: "IN PLAIN WORDS", 2: easy_first("competency::" + comp, 400)}, level=1, height=44, fl=fill("FFF8E1"), merge=(2, 11))
    row({1: "SKILLS FOR EVERY TOPIC", 2: "These apply whatever the child's needs — listed once here, not under each topic."},
        level=1, height=20, fl=fill("E3F2EE"), font=F(10, True, DARK), merge=(2, 11))
    M.cell(r, 2).font = F(9, False, MUTED, i=True)
    for s in every:
        skill_row(s, 2, True)
    if not specific:
        row({2: "Every %s skill applies to every topic, so there is no age-band or topic breakdown for this competency." % comp},
            level=1, height=20, font=F(9, False, MUTED, i=True), merge=(2, 11))
        continue
    for band in BANDS:
        br = row({1: "AGE BAND", 2: band}, level=1, height=22, fl=fill("2E5E55"), font=F(11, True, "FFFFFF"), merge=(2, 11))
        index.append(("   %s · %s" % (comp, band.split(" (")[0]), br))
        for topic in cond_order[band]:
            ts = [s for s in specific if topic in s["topics"]]
            if not ts:
                continue
            crow = cond_row[(band, topic)]
            applies = C.cell(crow, 7).value
            tr = row({1: "TOPIC", 2: topic, 3: str(applies or ""), 11: "full entry →"}, level=2, height=20, fl=fill("CFE8E1"),
                     font=F(10, True, DARK), links={2: cloc(crow), 11: cloc(crow)}, hidden=True)
            M.cell(tr, 3).font = F(8, False, MUTED, i=True); M.cell(tr, 11).font = F(8, False, LINK, u="single")
            # diagnosed
            row({1: "DIAGNOSED", 2: "%s  ·  %s" % (topic, C.cell(crow, 6).value or ""), 3: first_points(C.cell(crow, 9).value),
                 11: "full entry →"}, level=3, hidden=True, height=48, links={2: cloc(crow), 11: cloc(crow)})
            M.cell(r, 1).font = F(8, True, "8A5A00"); M.cell(r, 2).font = F(9, False, LINK, u="single")
            M.cell(r, 11).font = F(8, False, LINK, u="single")
            easy = easy_first("condition::" + topic)
            if easy:
                row({1: "IN PLAIN WORDS", 3: easy}, level=3, hidden=True, height=36, fl=fill("FFF8E1"))
                M.cell(r, 1).font = F(8, True, MUTED)
            if comp.startswith("1."):
                tl = tool_lines(C.cell(crow, 15).value)
                if tl:
                    row({1: "TOOLS AT THIS AGE", 3: tl}, level=3, hidden=True, height=30)
                    M.cell(r, 1).font = F(8, True, MUTED)
            # not diagnosed
            bp = band_pres_row.get((band, topic))
            plist = pres_by_topic.get(topic, [])
            if bp or plist:
                row({1: "NOT DIAGNOSED", 2: "Looks like %s, without the diagnosis" % topic.split(" (")[0],
                     11: "at this age →" if bp else None}, level=3, hidden=True, height=18,
                    links={11: cloc(bp) if bp else None}, font=F(9, True, "8A5A00"))
                if bp: M.cell(r, 11).font = F(8, False, LINK, u="single")
                for p in plist:
                    pr = pres_row.get(p)
                    row({2: p, 3: easy_first("presentation::" + p, 200) or first_points(C.cell(pr, 8).value if pr else "", 1, 200),
                         11: "full entry →" if pr else None}, level=3, hidden=True, height=30,
                        links={2: cloc(pr) if pr else None, 11: cloc(pr) if pr else None})
                    M.cell(r, 2).font = F(9, False, LINK, u="single"); M.cell(r, 3).font = F(8, False, INK)
                    M.cell(r, 11).font = F(8, False, LINK, u="single")
            # skills
            row({1: "SKILLS YOU NEED", 2: "Beginner → expert, in the order you learn them", 4: "1", 5: "2", 6: "3", 7: "4", 8: "5",
                 9: "Where I am now", 10: "When"}, level=3, hidden=True, height=18, fl=fill("E3F2EE"), font=F(8, True, DARK))
            for s in ts:
                skill_row(s, 3, True)

if unlinked:
    ur = row({1: "OTHER NEEDS — not linked to one diagnosis"}, height=26, fl=fill(DARK), font=F(13, True, "FFFFFF"), merge=(1, 11))
    index.append(("Other needs — not linked to one diagnosis", ur))
    for p in unlinked:
        pr = pres_row.get(p)
        row({2: p, 3: easy_first("presentation::" + p, 200), 11: "full entry →" if pr else None}, level=1, height=30,
            links={2: cloc(pr) if pr else None, 11: cloc(pr) if pr else None})
        M.cell(r, 2).font = F(9, False, LINK, u="single"); M.cell(r, 11).font = F(8, False, LINK, u="single")

# contents (in the rows kept free at the top)
r = index_at - 1
row({1: "CONTENTS — click to jump"}, height=20, font=F(11, True, DARK), merge=(1, 11))
for label, at in index:
    row({1: label}, height=15, font=F(9, not label.startswith(" "), LINK, u="single"), merge=(1, 6),
        links={1: "'%s'!A%d" % (NAME, at)})
assert r <= index_at + n_index, (r, index_at, n_index)
M.freeze_panes = "A2"

for ref, o, k in here_rules:
    M.conditional_formatting.add(ref,
                                 FormulaRule(formula=["IFERROR(VALUE(LEFT(%s,1)),0)>=%d" % (o, k + 1)],
                                             fill=fill(BAR[k]), font=Font(color="FFFFFF" if k >= 3 else INK, bold=True)))

wb.save(BOOK)
kinds = {}
for rr in range(1, M.max_row + 1):
    kinds[M.cell(rr, 1).value] = kinds.get(M.cell(rr, 1).value, 0) + 1
print("competency map: %d rows · topics %d · skills rows %d · presentations linked %d / unlinked %d"
      % (M.max_row, kinds.get("TOPIC", 0), kinds.get("SKILL", 0), sum(1 for v in PRES_TOPICS.values() if v), len(unlinked)))
