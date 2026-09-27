"""Corrections to text that was already in the workbook before this build.

Each fix is a (sheet, pattern, replacement, reason) rule applied to every cell of
that sheet. Runs in build_all.sh straight after src/base.xlsx is copied, so the
pristine base is never edited and every correction is visible here.
Writes CORRECTIONS.md listing what changed and how many cells each rule touched.
"""
import re
from openpyxl import load_workbook

BOOK = "The_Encyclopedia.xlsx"
TUSLA = (" This does not replace your own duty: if you are a mandated person (confirm your status with UCD), "
         "a report must reach Tusla as soon as practicable — jointly with the DLP or by you (Children First Act 2015).")

RULES = [
    # ── ICD-11 codes in Part D (repeated in every age band) ──────────────────
    ("Reference", r"(Stereotypic Movement Disorder — DSM-5-TR / ICD-11 )6A04", r"\g<1>6A06",
     "6A04 is Developmental motor coordination disorder; stereotyped movement disorder is 6A06."),
    ("Reference", r"(Reactive Attachment Disorder — DSM-5-TR / ICD-11 )6B40–6B4Z", r"\g<1>6B44",
     "6B40–6B4Z is the whole stress-associated block; reactive attachment disorder is 6B44."),
    ("Reference", r"(Adjustment Disorder \(with depressed mood / anxiety / mixed / conduct / unspecified\) — DSM-5-TR / ICD-11 )6B40–6B4Z",
     r"\g<1>6B43", "Adjustment disorder is 6B43, not the block range."),
    ("Reference", r"(Acute Stress Disorder — DSM-5-TR / ICD-11 )6B40–6B4Z",
     r"\g<1>QE84 (acute stress reaction — in ICD-11 not a mental disorder)",
     "ICD-11 codes acute stress reaction as QE84 (factors influencing health), not within 6B4x."),
    ("Reference", r"(Disruptive Mood Dysregulation Disorder — DSM-5-TR / ICD-11 )6A70–6A7Z",
     r"\g<1>— no ICD-11 equivalent (nearest: 6C90.0 ODD with chronic irritability-anger)",
     "DMDD is DSM-5-TR only; 6A70–6A7Z is the depressive-disorders block."),
    ("Reference", r"(Premenstrual Dysphoric Disorder — DSM-5-TR / ICD-11 )6A70–6A7Z", r"\g<1>GA34.41",
     "ICD-11 places PMDD at GA34.41 (with a cross-listing under depressive disorders)."),
    # ── safeguarding: the DLP route never discharges the mandated person's duty ──
    ("Reference", r"(you follow Children First and the school DLP route\.)", r"\1" + TUSLA,
     "The original sentence could be read as the DLP route being sufficient."),
    ("Reference", r"(and you follow the DLP route\.)", r"\1" + TUSLA, "As above."),
    ("Reference", r"(Follow the school's Designated Liaison Person route\.)", r"\1" + TUSLA, "As above."),
    ("Reference", r"(Tell your supervisor the same day and follow the school's DLP route\.)", r"\1" + TUSLA, "As above."),
    ("Reference", r"(Follow the school DLP route and NEPS protocol)(?! —)", r"\1 — and Tusla as soon as practicable if the threshold is met", "As above."),
    ("Reference", r"(Any concern goes straight to the supervisor and the DLP\.)",
     r"\1 A report to Tusla follows as soon as practicable where the threshold is met — the DLP does not discharge your own duty.",
     "As above."),
]

wb = load_workbook(BOOK)
log = ["# Corrections applied to the original workbook text", "",
       "Applied by fix_base.py on every build; src/base.xlsx is untouched.", "",
       "| Sheet | Rule | Cells changed | Reason |", "|---|---|---|---|"]
total = 0
for sheet, pat, rep, why in RULES:
    ws, rx, n = wb[sheet], re.compile(pat), 0
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and not c.value.startswith("=") and rx.search(c.value):
                c.value = rx.sub(rep, c.value); n += 1
    if n == 0:
        raise SystemExit("fix_base: rule matched nothing — %r" % pat)
    total += n
    log.append("| %s | `%s` | %d | %s |" % (sheet, pat.replace("|", "\\|")[:80], n, why))
wb.save(BOOK)
open("CORRECTIONS.md", "w").write("\n".join(log) + "\n")
print("fix_base: %d rules · %d cells corrected" % (len(RULES), total))
