"""Skill journey — every micro-skill as a 'slide' in a carousel, beginner to expert.

Order: the teaching timetable in Skill progression's BLOCK rows (Before placement →
Weeks 1–2 → Weeks 3–6 → Weeks 7–12 → Ongoing all year), then the workbook's own order.
Each slide reads left to right through the five stages (WATCH → LEAD AND TEACH), with
◀ PREVIOUS / NEXT ▶ links, the target dates, and a live YOU ARE HERE marker that follows
'Skill progression' column O. Contents at the top, with a live progress table.
Nothing is typed twice: text comes from Skill progression and records/easy_*.py.
"""
import glob, importlib.util, math, re
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter

BOOK = "The_Encyclopedia.xlsx"
NAME = "Skill journey"
SP = "'Skill progression'"

EASY = {}
for p in sorted(glob.glob("records/easy_*.py")):
    s = importlib.util.spec_from_file_location(p, p); m = importlib.util.module_from_spec(s); s.loader.exec_module(m)
    EASY.update(m.EASY)

PHASES = [
    ("Before placement", "BEFORE PLACEMENT", "Get ready. Paperwork, vetting, contracts and the ground rules — done once, before you meet a child."),
    ("Weeks 1–2", "WEEKS 1–2", "First steps. You watch, you learn the methods, you find your way around the service."),
    ("Weeks 3–6", "WEEKS 3–6", "Building up. You do more of the work alongside your supervisor, then begin to lead parts of it."),
    ("Weeks 7–12", "WEEKS 7–12", "Taking the lead. You run the work yourself and bring it to supervision."),
    ("Ongoing all year", "ALL YEAR", "Habits you keep. Skills you practise every week rather than learn once."),
]
STAGES = ["1 · WATCH", "2 · DO ALONGSIDE", "3 · LEAD, SUPERVISOR IN ROOM", "4 · LEAD, DISCUSS AFTER", "5 · LEAD AND TEACH"]
TAGS = ["BEGINNER", "", "", "", "EXPERT"]
HEAD = ["DCEFEA", "B7DFD5", "86C5B6", "4E9E8C", "1F6F5F"]   # light → dark: the progression
BODY = ["F4FAF8", "EAF5F2", "DDEFEA", "CFE8E1", "C0E0D7"]
INK, DARK, MUTED, HERE = "1B2B28", "123F37", "5B6B67", "FFD966"
CARD = [2, 4, 6, 8, 10]          # B D F H J — stage cards
ARROW = [3, 5, 7, 9]             # C E G I — arrows between them
fill = lambda h: PatternFill("solid", fgColor=h)
thin = Side(style="thin", color="9FBFB7")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = lambda h="left", v="top": Alignment(horizontal=h, vertical=v, wrap_text=True)

wb = load_workbook(BOOK)
S = wb["Skill progression"]
if NAME in wb.sheetnames:
    del wb[NAME]
J = wb.create_sheet(NAME, index=wb.sheetnames.index("Skill progression"))
J.sheet_view.showGridLines = False
J.sheet_properties.tabColor = "1F6F5F"
for col, w in {1: 2, 11: 2, **{c: 34 for c in CARD}, **{c: 4 for c in ARROW}}.items():
    J.column_dimensions[get_column_letter(col)].width = w

# ── collect skills in timetable order ────────────────────────────────────────
phase_of = lambda txt: next((i for i, (k, _, _) in enumerate(PHASES) if k in str(txt or "")), len(PHASES) - 1)
skills, cur = [], (len(PHASES) - 1, "")
for r in range(2, S.max_row + 1):
    kind = S.cell(r, 1).value
    if kind == "BLOCK":
        cur = (phase_of(S.cell(r, 6).value), str(S.cell(r, 7).value or S.cell(r, 5).value or ""))
    elif kind == "MICRO-SKILL":
        skills.append({"row": r, "phase": cur[0], "comp": S.cell(r, 2).value, "macro": S.cell(r, 3).value,
                       "block": S.cell(r, 4).value, "name": S.cell(r, 5).value, "std": S.cell(r, 6).value,
                       "why": S.cell(r, 7).value, "stages": [S.cell(r, c).value for c in range(8, 13)],
                       "error": S.cell(r, 13).value, "feeds": S.cell(r, 14).value, "next": S.cell(r, 16).value,
                       "targets": [None, None] + [S.cell(r, c).value for c in (18, 19, 20)], "evidence": S.cell(r, 21).value})
skills.sort(key=lambda s: (s["phase"], s["row"]))


def lines(text, width):
    """Rough wrapped-line count for 10pt text in a column span of `width`."""
    per = max(10, int(width * 1.15))
    return sum(max(1, math.ceil(len(p) / per)) for p in str(text or "").split("\n"))


def height(*pairs, lo=15, cap=409):
    return min(cap, max(lo, max(lines(t, w) for t, w in pairs) * 13.5 + 8))


def put(r, c, v, font=None, fl=None, al=None, border=None, merge_to=None, link=None):
    cell = J.cell(r, c); cell.value = v
    if font: cell.font = font
    if fl: cell.fill = fl
    cell.alignment = al or wrap()
    if border: cell.border = border
    if merge_to:
        J.merge_cells(start_row=r, start_column=c, end_row=r, end_column=merge_to)
        if fl:
            for k in range(c + 1, merge_to + 1): J.cell(r, k).fill = fl
    if link: cell.hyperlink = Hyperlink(ref=cell.coordinate, location="'%s'!B%d" % (NAME, link))
    return cell


# ── pass 1: work out where every slide starts (links need row numbers) ───────
SLIDE_ROWS = 10
n_phase = len({s["phase"] for s in skills})
toc_rows = 20 + n_phase + len(skills)                  # title block, progress table, contents lines
r, starts, phase_rows = toc_rows + 1, [], {}
last_phase = None
for s in skills:
    if s["phase"] != last_phase:
        phase_rows[s["phase"]] = r; r += 2; last_phase = s["phase"]
    starts.append(r); r += SLIDE_ROWS
by_name = {}
for i, s in enumerate(skills):
    by_name.setdefault(" ".join(str(s["name"]).split()), i)

# ── top: title, how to use, live progress, contents ─────────────────────────
put(1, 2, "SKILL JOURNEY — from beginner to expert", Font(name="Arial", size=18, bold=True, color="FFFFFF"), fill(DARK),
    wrap("left", "center"), merge_to=10)
J.row_dimensions[1].height = 34
put(2, 2, "⌂  Contents and progress  ·  each skill is one slide: use ◀ PREVIOUS and NEXT ▶ to move through them in order",
    Font(name="Arial", size=10, bold=True, color=DARK), fill("E3F2EE"), wrap("left", "center"), merge_to=10, link=4)
J.row_dimensions[2].height = 22
J.freeze_panes = "A3"
put(4, 2, "HOW IT WORKS\n"
          "• The skills are in the order you meet them: before placement, weeks 1–2, weeks 3–6, weeks 7–12, then habits you keep all year.\n"
          "• Each slide reads left to right, from BEGINNER (1 · Watch) to EXPERT (5 · Lead and teach).\n"
          "• The yellow card and ▲ YOU ARE HERE follow what you pick in Skill progression, column O. Click ✎ on a slide to go there and change it.\n"
          "• The plain-words version of every skill is at the bottom of its slide.",
    Font(name="Arial", size=10, color=INK), fill("F4FAF8"), merge_to=10, border=box)
J.row_dimensions[4].height = 80

# live progress table: competency × stage, counted from Skill progression
put(6, 2, "WHERE YOU ARE — counted live from Skill progression column O", Font(name="Arial", size=11, bold=True, color=DARK),
    merge_to=10, al=wrap("left", "center"))
put(7, 2, "Competency", Font(name="Arial", size=9, bold=True, color="FFFFFF"), fill(DARK), border=box, merge_to=3)
heads = [(4, "Not started", "9FB5B0")] + [(5 + i, STAGES[i].split(" · ")[0] + " " + ["Watch", "Alongside", "Lead + in room", "Lead + after", "Teach"][i], HEAD[i]) for i in range(5)]
for c, h, shade in heads:
    put(7, c, h, Font(name="Arial", size=9, bold=True, color="FFFFFF" if shade in HEAD[3:] or shade == "9FB5B0" else INK),
        fill(shade), wrap("center", "center"), border=box)
J.row_dimensions[7].height = 30
comps = list(dict.fromkeys(s["comp"] for s in skills))
for i, comp in enumerate(comps):
    rr = 8 + i
    put(rr, 2, comp, Font(name="Arial", size=9, color=INK), border=box, merge_to=3, al=wrap("left", "center"))
    for c, key in zip(range(4, 10), ["Not started", "1*", "2*", "3*", "4*", "5*"]):
        put(rr, c, '=COUNTIFS(%s!$A:$A,"MICRO-SKILL",%s!$B:$B,$B%d,%s!$O:$O,"%s")' % (SP, SP, rr, SP, key),
            Font(name="Arial", size=9, color=INK), al=wrap("center", "center"), border=box)
put(16, 2, "Stage 5 means you can teach it. Aim to move one stage at a time.", Font(name="Arial", size=9, italic=True, color=MUTED), merge_to=10)

# contents list
rr = 18
put(rr, 2, "CONTENTS — click a skill to open its slide", Font(name="Arial", size=11, bold=True, color=DARK), merge_to=10); rr += 1
last_phase = None
for i, s in enumerate(skills):
    if s["phase"] != last_phase:
        last_phase = s["phase"]
        put(rr, 2, "%s — %s" % (PHASES[s["phase"]][1], PHASES[s["phase"]][2]), Font(name="Arial", size=10, bold=True, color="FFFFFF"),
            fill(HEAD[min(4, s["phase"])] if s["phase"] >= 3 else DARK), merge_to=10, al=wrap("left", "center"), link=phase_rows[s["phase"]])
        J.row_dimensions[rr].height = 18; rr += 1
    put(rr, 2, "%3d.  %s" % (i + 1, s["name"]), Font(name="Arial", size=9, color="1F4E79", underline="single"),
        merge_to=7, al=wrap("left", "center"), link=starts[i])
    put(rr, 8, s["comp"], Font(name="Arial", size=8, color=MUTED), al=wrap("left", "center"))
    put(rr, 10, "=%s!O%d" % (SP, s["row"]), Font(name="Arial", size=8, color=MUTED), al=wrap("left", "center"))
    rr += 1
assert rr <= toc_rows + 1, (rr, toc_rows)

# ── slides ───────────────────────────────────────────────────────────────────
here_rules = []
for i, s in enumerate(skills):
    if s["phase"] in phase_rows and phase_rows[s["phase"]] == starts[i] - 2:
        pr = starts[i] - 2
        p = PHASES[s["phase"]]
        put(pr, 2, "%s\n%s" % (p[1], p[2]), Font(name="Arial", size=14, bold=True, color="FFFFFF"), fill(DARK),
            wrap("left", "center"), merge_to=10)
        J.row_dimensions[pr].height = 48
        J.row_dimensions[pr + 1].height = 8
    r0, o = starts[i], "%s!$O$%d" % (SP, s["row"])
    # r0 · navigation bar
    nav_f, nav_fill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), fill("2E5E55")
    put(r0, 2, "◀  PREVIOUS" if i else "⌂  CONTENTS", nav_f, nav_fill, wrap("left", "center"), link=starts[i - 1] if i else 4)
    put(r0, 3, None, fl=nav_fill)
    put(r0, 4, "SKILL %d OF %d  ·  %s  ·  %s" % (i + 1, len(skills), PHASES[s["phase"]][1], s["comp"]), nav_f, nav_fill,
        wrap("center", "center"), merge_to=8)
    put(r0, 9, None, fl=nav_fill)
    put(r0, 10, "NEXT  ▶" if i < len(skills) - 1 else "⌂  CONTENTS", nav_f, nav_fill, wrap("right", "center"),
        link=starts[i + 1] if i < len(skills) - 1 else 4)
    J.row_dimensions[r0].height = 22
    # r0+1 · title
    put(r0 + 1, 2, s["name"], Font(name="Arial", size=15, bold=True, color=DARK), merge_to=10, al=wrap("left", "center"))
    J.row_dimensions[r0 + 1].height = height((s["name"], 140), lo=28) * 1.4
    # r0+2 · where it sits
    put(r0 + 2, 2, "%s  ›  %s  ·  %s" % (s["macro"], s["block"], s["std"] or ""), Font(name="Arial", size=9, italic=True, color=MUTED),
        merge_to=10, al=wrap("left", "center"))
    J.row_dimensions[r0 + 2].height = 16
    # r0+3 · why
    put(r0 + 3, 2, "WHY IT MATTERS\n" + str(s["why"] or ""), Font(name="Arial", size=10, color=INK), fill("F7F7F2"), merge_to=10, border=box)
    J.row_dimensions[r0 + 3].height = height(("WHY IT MATTERS\n" + str(s["why"] or ""), 180))
    # r0+4 · stage headers (the progression)
    for k, c in enumerate(CARD):
        put(r0 + 4, c, STAGES[k] + ("\n" + TAGS[k] if TAGS[k] else ""), Font(name="Arial", size=10, bold=True,
            color="FFFFFF" if k >= 3 else INK), fill(HEAD[k]), wrap("center", "center"), border=box)
    for c in ARROW:
        put(r0 + 4, c, "➜", Font(name="Arial", size=14, bold=True, color="4E9E8C"), al=wrap("center", "center"))
        put(r0 + 5, c, "➜", Font(name="Arial", size=14, bold=True, color="B7DFD5"), al=wrap("center", "center"))
    J.row_dimensions[r0 + 4].height = 32
    # r0+5 · stage cards
    for k, c in enumerate(CARD):
        put(r0 + 5, c, s["stages"][k], Font(name="Arial", size=10, color=INK), fill(BODY[k]), border=box)
        L = get_column_letter(c)
        here_rules.append(("%s%d:%s%d" % (L, r0 + 5, L, r0 + 6), o, k + 1))
    J.row_dimensions[r0 + 5].height = height(*[(t, 34) for t in s["stages"]])
    # r0+6 · you are here + target dates
    for k, c in enumerate(CARD):
        tgt = s["targets"][k]
        put(r0 + 6, c, '=IF(LEFT(%s,1)="%d","▲ YOU ARE HERE","")%s' % (o, k + 1, ('&"   Aim: %s"' % tgt) if tgt else ""),
            Font(name="Arial", size=9, bold=True, color=DARK), al=wrap("center", "center"))
    J.row_dimensions[r0 + 6].height = 18
    # r0+7 · error · feeds into · next step
    feeds = str(s["feeds"] or "")
    put(r0 + 7, 2, "THE ERROR THAT COSTS MARKS\n" + str(s["error"] or ""), Font(name="Arial", size=9, color=INK), fill("FBEDEA"), merge_to=4, border=box)
    hit = min(((feeds.find(nm), nm) for nm in by_name if len(nm) > 12 and nm in feeds and nm != s["name"]), default=None)
    put(r0 + 7, 5, "WHAT IT LEADS TO NEXT" + ("  (click to jump)" if hit else "") + "\n" + feeds, Font(name="Arial", size=9, color=INK),
        fill("EAF1FB"), merge_to=7, border=box, link=starts[by_name[hit[1]]] if hit else None)
    nxt = ('="YOUR ONE NEXT STEP — the stage after where you are now"&CHAR(10)&IF(LEFT(%s,1)="5",'
           '"You can teach this skill. Keep it going, and help someone else learn it.",'
           'INDEX(%s!$H$%d:$L$%d,1,IFERROR(VALUE(LEFT(%s,1)),0)+1))' % (o, SP, s["row"], s["row"], o))
    put(r0 + 7, 8, nxt, Font(name="Arial", size=9, color=INK), fill("EEF6E8"), merge_to=10, border=box)
    J.row_dimensions[r0 + 7].height = height(("THE ERROR THAT COSTS MARKS\n" + str(s["error"] or ""), 72), ("x\n" + feeds, 72),
                                             *[("x\n" + str(t or ""), 72) for t in s["stages"]])
    # r0+8 · where I am (live) · evidence
    c = put(r0 + 8, 2, '="WHERE I AM NOW:  "&%s&"      ✎ change it"' % o, Font(name="Arial", size=9, bold=True, color="1F4E79", underline="single"),
            merge_to=5, al=wrap("left", "center"))
    c.hyperlink = Hyperlink(ref=c.coordinate, location="%s!O%d" % (SP, s["row"]))
    put(r0 + 8, 6, "EVIDENCE GOES IN:  " + str(s["evidence"] or ""), Font(name="Arial", size=9, color=MUTED), merge_to=10, al=wrap("left", "center"))
    J.row_dimensions[r0 + 8].height = height(("EVIDENCE GOES IN:  " + str(s["evidence"] or ""), 110), lo=18)
    # r0+9 · easy read
    easy = EASY.get("micro::" + str(s["name"]))
    if easy:
        put(r0 + 9, 2, "IN PLAIN WORDS\n" + easy, Font(name="Arial", size=10, color=INK), fill("FFF8E1"), merge_to=10, border=box)
        J.row_dimensions[r0 + 9].height = height(("IN PLAIN WORDS\n" + easy, 180))
    else:
        J.row_dimensions[r0 + 9].height = 6

# highlight the card for the stage you are at
for ref, o, k in here_rules:
    J.conditional_formatting.add(ref, FormulaRule(formula=['LEFT(%s,1)="%d"' % (o, k)], fill=fill(HERE),
                                                  font=Font(bold=True, color=INK)))

wb.save(BOOK)
print("skill journey: %d slides · %s" % (len(skills), " · ".join(
    "%s %d" % (PHASES[p][1], sum(1 for s in skills if s["phase"] == p)) for p in range(len(PHASES)))))
